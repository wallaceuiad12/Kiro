# -*- coding: utf-8 -*-
"""
Insere as respostas da lista de Economia Política II dentro do próprio .docx das
questões, logo abaixo de cada enunciado, e gera também uma versão em Markdown.

Uso:
    python3 ce405-ep2/gerar.py

Entrada : ce_405_c_lista_1_exercicios.docx  (original, versionado no git)
Saídas  : ce_405_c_lista_1_exercicios.docx  (reescrito com as respostas)
          ce405-ep2/lista1-resolvida.md     (mesmo conteúdo em Markdown)

O original permanece recuperável pelo git (git checkout -- <arquivo>), por isso a
escrita é feita no próprio arquivo, como pede o enunciado da tarefa.
"""
import os
import re
import shutil
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from respostas import RESPOSTAS  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AQUI = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(ROOT, 'ce_405_c_lista_1_exercicios.docx')
ORIGINAL = os.path.join(AQUI, 'ce_405_c_lista_1_exercicios-original.docx')
MD = os.path.join(AQUI, 'lista1-resolvida.md')

FONTE = 'Times New Roman'
AZUL = RGBColor(0x1F, 0x38, 0x64)
AZUL2 = RGBColor(0x1F, 0x4E, 0x79)
CINZA = RGBColor(0x40, 0x40, 0x40)
SOMBRA_CAB = 'DCE6F1'
SOMBRA_FORM = 'F2F2F2'

TOKEN = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')


# --------------------------------------------------------------------- helpers
def _xml(tag, **attrs):
    from docx.oxml import OxmlElement
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn('w:' + k), v)
    return el


# O esquema do OOXML exige ordem fixa dos filhos de pPr, tblPr e tcPr. Anexar no
# fim gera arquivo que o Word acusa como corrompido, por isso cada elemento é
# inserido antes dos seus sucessores legítimos (listas extraídas de CT_PPr,
# CT_TblPr e CT_TcPr do python-docx).
_SUC_PBDR = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
             'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
             'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
             'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
             'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
             'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
_SUC_PSHD = _SUC_PBDR[1:]
_SUC_TBLBORDERS = ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
                   'w:tblDescription', 'w:tblPrChange')
_SUC_TCSHD = ('w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
              'w:hideMark', 'w:headers', 'w:cellIns', 'w:cellDel', 'w:cellMerge',
              'w:tcPrChange')


def _sombrear_par(par, cor):
    pPr = par._p.get_or_add_pPr()
    pPr.insert_element_before(_xml('w:shd', val='clear', color='auto', fill=cor),
                              *_SUC_PSHD)


def _sombrear_celula(celula, cor):
    tcPr = celula._tc.get_or_add_tcPr()
    tcPr.insert_element_before(_xml('w:shd', val='clear', color='auto', fill=cor),
                               *_SUC_TCSHD)


def _borda_superior(par, cor='BFBFBF', tamanho='8'):
    pPr = par._p.get_or_add_pPr()
    bordas = _xml('w:pBdr')
    bordas.append(_xml('w:top', val='single', sz=tamanho, space='6', color=cor))
    pPr.insert_element_before(bordas, *_SUC_PBDR)


def _bordas_tabela(tabela, cor='9CB3C9'):
    tblPr = tabela._tbl.tblPr
    bordas = _xml('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        bordas.append(_xml('w:' + lado, val='single', sz='6', space='0', color=cor))
    tblPr.insert_element_before(bordas, *_SUC_TBLBORDERS)


def _runs(par, texto, tamanho=11.5, cor=None, negrito_tudo=False):
    """Escreve `texto` no parágrafo, interpretando **negrito** e *itálico*."""
    for pedaco in TOKEN.split(texto):
        if not pedaco:
            continue
        negrito, italico = negrito_tudo, False
        if pedaco.startswith('**') and pedaco.endswith('**'):
            pedaco, negrito = pedaco[2:-2], True
        elif pedaco.startswith('*') and pedaco.endswith('*'):
            pedaco, italico = pedaco[1:-1], True
        run = par.add_run(pedaco)
        run.font.name = FONTE
        run.font.size = Pt(tamanho)
        run.font.bold = negrito
        run.font.italic = italico
        if cor is not None:
            run.font.color.rgb = cor
    return par


def _novo_par(doc, esquerda=0.6, antes=0, depois=4, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
              pendente=None):
    par = doc.add_paragraph()
    par.style = doc.styles['Normal']
    pf = par.paragraph_format
    pf.left_indent = Cm(esquerda)
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.alignment = alinhamento
    pf.line_spacing = 1.12
    if pendente is not None:
        pf.first_line_indent = Cm(-pendente)
    # garante que o parágrafo não herde numeração automática da lista do documento
    pPr = par._p.get_or_add_pPr()
    for numPr in pPr.findall(qn('w:numPr')):
        pPr.remove(numPr)
    return par


# ------------------------------------------------------------------ renderizador
def construir_blocos(doc, resposta):
    """Cria os elementos XML de uma resposta e devolve-os na ordem de leitura."""
    elementos = []

    cab = _novo_par(doc, esquerda=0.5, antes=12, depois=5,
                    alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    _borda_superior(cab)
    _runs(cab, resposta['rotulo'], tamanho=11.5, cor=AZUL, negrito_tudo=True)
    elementos.append(cab._p)

    for tipo, conteudo in resposta['blocos']:
        if tipo == 'h':
            par = _novo_par(doc, esquerda=0.6, antes=8, depois=2,
                            alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
            _runs(par, conteudo, tamanho=11, cor=AZUL2, negrito_tudo=True)
            elementos.append(par._p)

        elif tipo == 'p':
            par = _novo_par(doc, esquerda=0.6, depois=5)
            _runs(par, conteudo, tamanho=11.5)
            elementos.append(par._p)

        elif tipo == 'b':
            par = _novo_par(doc, esquerda=1.1, depois=3, pendente=0.5)
            _runs(par, '• ' + conteudo, tamanho=11.5)
            elementos.append(par._p)

        elif tipo == 'd':
            rotulo, _, corpo = conteudo.partition('||')
            par = _novo_par(doc, esquerda=1.1, depois=3, pendente=0.5)
            _runs(par, '• ', tamanho=11.5)
            _runs(par, rotulo + '. ', tamanho=11.5, negrito_tudo=True)
            _runs(par, corpo, tamanho=11.5)
            elementos.append(par._p)

        elif tipo == 'f':
            par = _novo_par(doc, esquerda=1.3, antes=4, depois=6,
                            alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
            _sombrear_par(par, SOMBRA_FORM)
            run = _runs(par, conteudo, tamanho=12, negrito_tudo=True).runs[0]
            run.font.name = FONTE
            elementos.append(par._p)

        elif tipo == 't':
            linhas = conteudo
            tabela = doc.add_table(rows=len(linhas), cols=len(linhas[0]))
            tabela.style = doc.styles['Normal Table']
            tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
            tabela.autofit = True
            _bordas_tabela(tabela)
            for i, linha in enumerate(linhas):
                for j, celula_txt in enumerate(linha):
                    celula = tabela.cell(i, j)
                    celula.text = ''
                    par = celula.paragraphs[0]
                    pf = par.paragraph_format
                    pf.space_before, pf.space_after = Pt(2), Pt(2)
                    pf.line_spacing = 1.0
                    _runs(par, celula_txt, tamanho=10, negrito_tudo=(i == 0))
                    if i == 0:
                        _sombrear_celula(celula, SOMBRA_CAB)
            elementos.append(tabela._tbl)
            # Word exige um parágrafo depois da tabela para não colar no texto seguinte
            elementos.append(_novo_par(doc, esquerda=0.6, antes=0, depois=2)._p)

        else:
            raise ValueError('tipo de bloco desconhecido: %r' % tipo)

    return elementos


def nota_de_cabecalho(doc):
    par = _novo_par(doc, esquerda=0, antes=2, depois=10,
                    alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    _runs(par, 'Respostas completas inseridas abaixo de cada questão. '
               'Referência: MARX, K. O Capital, Livro I (ed. Boitempo).',
          tamanho=10, cor=CINZA)
    for run in par.runs:
        run.font.italic = True
    return par._p


# ------------------------------------------------------------------------ main
def escrever_docx():
    if not os.path.exists(ORIGINAL):
        shutil.copy2(DOCX, ORIGINAL)   # guarda uma cópia limpa das questões
    doc = docx.Document(ORIGINAL)      # parte sempre do original, idempotente
    paragrafos = list(doc.paragraphs)

    for resposta in RESPOSTAS:
        indice = resposta['depois']
        ancora = paragrafos[indice]._p
        for elemento in reversed(construir_blocos(doc, resposta)):
            ancora.addnext(elemento)

    paragrafos[0]._p.addnext(nota_de_cabecalho(doc))
    doc.save(DOCX)
    return DOCX


def escrever_markdown():
    doc = docx.Document(ORIGINAL)
    paragrafos = [p.text.strip() for p in doc.paragraphs]
    por_indice = {r['depois']: r for r in RESPOSTAS}

    saida = ['# Lista de Questões I de Economia Política II (CE405, turma C, 2025)',
             '',
             'Questões transcritas do enunciado original, com as respostas logo abaixo de '
             'cada uma. Referência: MARX, Karl. *O Capital*, Livro I (ed. Boitempo).',
             '']

    for i, texto in enumerate(paragrafos):
        if texto:
            if texto.lower().startswith('capítulo'):
                saida += ['', '## ' + texto, '']
            elif i == 0:
                pass
            else:
                saida += ['> ' + texto, '']
        if i in por_indice:
            r = por_indice[i]
            saida += ['### ' + r['rotulo'], '']
            for tipo, conteudo in r['blocos']:
                if tipo == 'h':
                    if saida and saida[-1] != '':
                        saida.append('')          # separa o subtítulo da lista anterior
                    saida += ['**' + conteudo.replace('**', '') + '**', '']
                elif tipo == 'p':
                    saida += [conteudo, '']
                elif tipo == 'b':
                    saida.append('- ' + conteudo)
                elif tipo == 'd':
                    rotulo, _, corpo = conteudo.partition('||')
                    saida.append('- **' + rotulo + '.** ' + corpo)
                elif tipo == 'f':
                    saida += ['', '```', conteudo, '```', '']
                elif tipo == 't':
                    linhas = conteudo
                    saida.append('')
                    saida.append('| ' + ' | '.join(linhas[0]) + ' |')
                    saida.append('|' + '---|' * len(linhas[0]))
                    for linha in linhas[1:]:
                        saida.append('| ' + ' | '.join(linha) + ' |')
                    saida.append('')
            saida.append('')

    with open(MD, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(saida).replace('\n\n\n', '\n\n') + '\n')
    return MD


if __name__ == '__main__':
    print('respostas carregadas:', len(RESPOSTAS))
    print('docx     ->', escrever_docx())
    print('markdown ->', escrever_markdown())
